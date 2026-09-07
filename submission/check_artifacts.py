"""Check submission contracts after export; page images still require visual review."""

import json
import os
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from pypdf import PdfReader

out = Path(os.environ["PAP_SUBMISSION_OUT"])
source = Path(os.environ["PAP_SUBMISSION_SRC"])
content = json.loads((source / "content.json").read_text())
report = Document(out / "DragonWings_Final_Report.docx")
texts = [p.text for p in report.paragraphs]
texts += [cell.text for table in report.tables for row in table.rows for cell in row.cells]
word_count = len(re.findall(r"\S+", " ".join(texts)))
assert 1000 <= word_count <= 1500, f"Report outside suggested word range: {word_count}"
assert all(section["heading"] in texts for page in content["report"]["pages"] for section in page)
assert 480 <= sum(s["seconds"] for s in content["slides"]) <= 600
assert sum(s["seconds"] for s in content["pitch_slides"]) == 90
ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
for name, slides in [
    ("DragonWings_Final_Presentation", content["slides"]),
    ("DragonWings_90_Second_Pitch", content["pitch_slides"]),
]:
    with zipfile.ZipFile(out / (name + ".pptx")) as package:
        slide_files = [
            n for n in package.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)
        ]
        note_files = [
            n for n in package.namelist() if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", n)
        ]
        assert len(slide_files) == len(slides)
        assert len(note_files) == len(slides)
        for index, slide in enumerate(slides, 1):
            root = ET.fromstring(package.read(f"ppt/notesSlides/notesSlide{index}.xml"))
            note = " ".join(x.text or "" for x in root.findall(".//a:t", ns))
            assert slide["notes"].strip(), f"Empty narration for slide {index}"
            expected = " ".join(slide["notes"].split())
            assert expected in " ".join(note.split()), f"Narration mismatch on slide {index}"
            assert f"Planned duration: {slide['seconds']} seconds" in note
        for filename in slide_files + note_files:
            text = " ".join(
                x.text or "" for x in ET.fromstring(package.read(filename)).findall(".//a:t", ns)
            )
            assert "{{" not in text, f"Unexpanded content in {filename}"
        # Each slide must have native editable text, not just a flattened image.
        for filename in slide_files:
            assert ET.fromstring(package.read(filename)).findall(".//a:t", ns), filename
    assert len(PdfReader(out / (name + ".pdf")).pages) == len(slides)

for docx in out.glob("*.docx"):
    with zipfile.ZipFile(docx) as package:
        assert b"{{" not in package.read("word/document.xml"), docx.name
        assert b"<w:pBdr" not in package.read("word/styles.xml"), docx.name
assert len(PdfReader(out / "DragonWings_Final_Report.pdf").pages) >= 1
receipt = {
    "report_words": word_count,
    "main_slides": len(content["slides"]),
    "main_seconds": sum(s["seconds"] for s in content["slides"]),
    "pitch_slides": len(content["pitch_slides"]),
    "pitch_seconds": 90,
    "video_link_present": bool(content["video_url"]),
    "visual_review": "Required separately for all rendered pages and slides",
}
Path(os.environ["PAP_SUBMISSION_BUILD"], "content-checks.json").write_text(
    json.dumps(receipt, indent=2) + "\n"
)
print(json.dumps(receipt, indent=2))
